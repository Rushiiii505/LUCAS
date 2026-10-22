import csv
import json
import io
from pathlib import Path
from typing import List, Dict, Any, Optional
from pypdf import PdfReader
import docx
from app.utils.logger import logger

class ParsedSection:
    def __init__(self, content: str, page_or_section: Optional[str] = None):
        self.content = content.strip()
        self.page_or_section = page_or_section

def clean_whitespace(text: str) -> str:
    """Normalize excess whitespace and line breaks while keeping paragraph flow."""
    lines = [line.strip() for line in text.splitlines()]
    clean_text = "\n".join(lines)
    # Collapse 3+ newlines into 2
    import re
    return re.sub(r'\n{3,}', '\n\n', clean_text).strip()

def parse_pdf(file_path: Path) -> List[ParsedSection]:
    sections: List[ParsedSection] = []
    try:
        reader = PdfReader(str(file_path))
        if len(reader.pages) == 0:
            raise ValueError("PDF contains no pages")

        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            cleaned = clean_whitespace(text)
            if cleaned:
                sections.append(ParsedSection(content=cleaned, page_or_section=f"Page {idx + 1}"))
    except Exception as e:
        logger.error(f"Error parsing PDF {file_path}: {e}")
        raise ValueError(f"Failed to parse PDF: {str(e)}")

    if not sections:
        raise ValueError("No readable text found in PDF document")
    return sections

def parse_docx(file_path: Path) -> List[ParsedSection]:
    sections: List[ParsedSection] = []
    try:
        doc = docx.Document(str(file_path))
        current_section = "Document Header"
        current_paragraphs: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            # If it's a heading
            if p.style.name.startswith("Heading"):
                if current_paragraphs:
                    combined = clean_whitespace("\n".join(current_paragraphs))
                    if combined:
                        sections.append(ParsedSection(content=combined, page_or_section=current_section))
                    current_paragraphs = []
                current_section = f"Section: {text}"
            else:
                current_paragraphs.append(text)

        if current_paragraphs:
            combined = clean_whitespace("\n".join(current_paragraphs))
            if combined:
                sections.append(ParsedSection(content=combined, page_or_section=current_section))

        # Also extract tables
        for t_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    table_rows.append(" | ".join(row_cells))
            if table_rows:
                t_content = "\n".join(table_rows)
                sections.append(ParsedSection(content=t_content, page_or_section=f"Table {t_idx + 1}"))

    except Exception as e:
        logger.error(f"Error parsing DOCX {file_path}: {e}")
        raise ValueError(f"Failed to parse DOCX: {str(e)}")

    if not sections:
        raise ValueError("No text content could be extracted from DOCX")
    return sections

def parse_markdown(file_path: Path) -> List[ParsedSection]:
    sections: List[ParsedSection] = []
    try:
        text = file_path.read_text(encoding="utf-8", errors="replace")
        lines = text.splitlines()
        current_header = "Preamble"
        current_block: List[str] = []

        for line in lines:
            if line.startswith("#"):
                if current_block:
                    content = clean_whitespace("\n".join(current_block))
                    if content:
                        sections.append(ParsedSection(content=content, page_or_section=current_header))
                    current_block = []
                current_header = line.lstrip("#").strip()
            else:
                current_block.append(line)

        if current_block:
            content = clean_whitespace("\n".join(current_block))
            if content:
                sections.append(ParsedSection(content=content, page_or_section=current_header))

    except Exception as e:
        logger.error(f"Error parsing Markdown {file_path}: {e}")
        raise ValueError(f"Failed to read Markdown file: {str(e)}")

    if not sections:
        raise ValueError("Markdown file is empty")
    return sections

def parse_text(file_path: Path) -> List[ParsedSection]:
    try:
        content = file_path.read_text(encoding="utf-8", errors="replace")
        cleaned = clean_whitespace(content)
        if not cleaned:
            raise ValueError("File is empty")
        return [ParsedSection(content=cleaned, page_or_section="Full Text")]
    except Exception as e:
        logger.error(f"Error parsing text file {file_path}: {e}")
        raise ValueError(f"Failed to read text file: {str(e)}")

def parse_csv(file_path: Path) -> List[ParsedSection]:
    sections: List[ParsedSection] = []
    try:
        with open(file_path, mode="r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            header = next(reader, None)
            if not header:
                raise ValueError("CSV file is empty")
            
            header_str = ", ".join(header)
            batch: List[str] = []
            row_count = 0

            for row in reader:
                if not any(cell.strip() for cell in row):
                    continue
                row_count += 1
                row_items = [f"{header[i]}: {row[i]}" for i in range(min(len(header), len(row))) if row[i].strip()]
                batch.append("; ".join(row_items))
                
                # Group 25 rows per section
                if len(batch) >= 25:
                    sections.append(ParsedSection(
                        content="\n".join(batch),
                        page_or_section=f"Rows {row_count - len(batch) + 1}–{row_count}"
                    ))
                    batch = []

            if batch:
                sections.append(ParsedSection(
                    content="\n".join(batch),
                    page_or_section=f"Rows {row_count - len(batch) + 1}–{row_count}"
                ))

    except Exception as e:
        logger.error(f"Error parsing CSV {file_path}: {e}")
        raise ValueError(f"Failed to parse CSV: {str(e)}")

    if not sections:
        raise ValueError("No data rows found in CSV")
    return sections

def parse_json(file_path: Path) -> List[ParsedSection]:
    sections: List[ParsedSection] = []
    try:
        content = file_path.read_text(encoding="utf-8", errors="replace")
        data = json.loads(content)
        
        if isinstance(data, list):
            for idx, item in enumerate(data):
                dumped = json.dumps(item, indent=2)
                sections.append(ParsedSection(content=dumped, page_or_section=f"Item {idx + 1}"))
        elif isinstance(data, dict):
            for key, val in data.items():
                val_str = json.dumps(val, indent=2) if isinstance(val, (dict, list)) else str(val)
                sections.append(ParsedSection(content=f"{key}:\n{val_str}", page_or_section=f"Key: {key}"))
        else:
            sections.append(ParsedSection(content=str(data), page_or_section="Data"))

    except Exception as e:
        logger.error(f"Error parsing JSON {file_path}: {e}")
        raise ValueError(f"Failed to parse JSON: {str(e)}")

    if not sections:
        raise ValueError("JSON file contains no parseable content")
    return sections

def parse_document(file_path: Path, file_type: str) -> List[ParsedSection]:
    """Unified entry point for document parsing based on file extension."""
    ft = file_type.lower().lstrip(".")
    if ft == "pdf":
        return parse_pdf(file_path)
    elif ft in ("docx", "doc"):
        return parse_docx(file_path)
    elif ft in ("md", "markdown"):
        return parse_markdown(file_path)
    elif ft == "csv":
        return parse_csv(file_path)
    elif ft == "json":
        return parse_json(file_path)
    elif ft in ("txt", "log", "text"):
        return parse_text(file_path)
    else:
        raise ValueError(f"Unsupported document format: {file_type}")
